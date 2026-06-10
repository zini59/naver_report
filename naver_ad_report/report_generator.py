"""리포트 시트 생성 (광고그룹별, 일자별, 요일별 등)"""

from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import KOREAN_WEEKDAYS
from .report_data import ReportBundle


def _add_weekday(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "일자" in out.columns and out["일자"].notna().any():
        out["요일"] = out["일자"].dt.dayofweek.map(lambda i: KOREAN_WEEKDAYS[i] if pd.notna(i) else "")
    else:
        out["요일"] = ""
    return out


def _calc_metrics_row(imp: float, clk: float, cost: float, conv: float = 0) -> dict:
    """최종 리포트 형식 지표 (클릭률=소수)."""
    ctr = round(clk / imp, 6) if imp > 0 else 0
    cpc = round(cost / clk, 2) if clk > 0 else 0
    cvr = round(conv / clk, 6) if clk > 0 else 0
    cpa = round(cost / conv, 2) if conv > 0 else 0
    return {
        "노출수": int(imp),
        "클릭수": int(clk),
        "클릭률": ctr,
        "CPC": cpc,
        "광고비": int(cost),
        "전환수": int(conv),
        "전환율": cvr,
        "CPA": cpa,
    }


def _aggregate(df: pd.DataFrame, group_cols: list[str], report_style: bool = False) -> pd.DataFrame:
    """노출/클릭/비용 합산 후 파생 지표 재계산."""
    if df.empty:
        return pd.DataFrame()

    work = df.copy()
    for col in ["노출수", "클릭수", "광고비", "전환수"]:
        if col not in work.columns:
            work[col] = 0

    if not group_cols:
        m = _calc_metrics_row(work["노출수"].sum(), work["클릭수"].sum(), work["광고비"].sum(), work["전환수"].sum())
        return pd.DataFrame([m])

    grouped = work.groupby(group_cols, dropna=False, as_index=False).agg(
        노출수=("노출수", "sum"),
        클릭수=("클릭수", "sum"),
        광고비=("광고비", "sum"),
        전환수=("전환수", "sum"),
    )

    if report_style:
        rows = []
        for _, r in grouped.iterrows():
            base = {c: r[c] for c in group_cols}
            base.update(_calc_metrics_row(r["노출수"], r["클릭수"], r["광고비"], r["전환수"]))
            rows.append(base)
        return pd.DataFrame(rows)

    grouped["클릭률(%)"] = grouped.apply(
        lambda r: round(r["클릭수"] / r["노출수"] * 100, 2) if r["노출수"] > 0 else 0,
        axis=1,
    )
    grouped["평균CPC"] = grouped.apply(
        lambda r: round(r["광고비"] / r["클릭수"], 0) if r["클릭수"] > 0 else 0,
        axis=1,
    )
    grouped["전환율(%)"] = grouped.apply(
        lambda r: round(r["전환수"] / r["클릭수"] * 100, 2) if r["클릭수"] > 0 else 0,
        axis=1,
    )
    grouped["전환당비용"] = grouped.apply(
        lambda r: round(r["광고비"] / r["전환수"], 0) if r["전환수"] > 0 else 0,
        axis=1,
    )
    return grouped


def build_summary_overview(df: pd.DataFrame) -> pd.DataFrame:
    """전체 요약."""
    if df.empty:
        return pd.DataFrame([{"항목": "데이터 없음"}])

    total_imp = df["노출수"].sum()
    total_clk = df["클릭수"].sum()
    total_cost = df["광고비"].sum()
    total_conv = df["전환수"].sum()

    date_range = ""
    if df["일자"].notna().any():
        dmin = df["일자"].min()
        dmax = df["일자"].max()
        date_range = f"{dmin.strftime('%Y-%m-%d')} ~ {dmax.strftime('%Y-%m-%d')}"

    rows = [
        {"항목": "분석기간", "값": date_range},
        {"항목": "총 노출수", "값": int(total_imp)},
        {"항목": "총 클릭수", "값": int(total_clk)},
        {"항목": "총 광고비", "값": int(total_cost)},
        {"항목": "총 전환수", "값": int(total_conv)},
        {"항목": "평균 클릭률(%)", "값": round(total_clk / total_imp * 100, 2) if total_imp else 0},
        {"항목": "평균 CPC", "값": round(total_cost / total_clk, 0) if total_clk else 0},
        {"항목": "평균 전환율(%)", "값": round(total_conv / total_clk * 100, 2) if total_clk else 0},
        {"항목": "전환당비용", "값": round(total_cost / total_conv, 0) if total_conv else 0},
        {"항목": "광고그룹 수", "값": df["광고그룹"].nunique() if "광고그룹" in df.columns else 0},
        {"항목": "캠페인 수", "값": df["캠페인"].nunique() if "캠페인" in df.columns else 0},
    ]
    return pd.DataFrame(rows)


def build_keyword_report(df: pd.DataFrame) -> pd.DataFrame:
    """키워드별 상세 (클릭순)."""
    if df.empty or "키워드" not in df.columns:
        return pd.DataFrame()

    group_cols = [c for c in ["캠페인", "광고그룹", "키워드", "검색유형"] if c in df.columns]
    result = _aggregate(df, group_cols, report_style=True)
    return result.sort_values("클릭수", ascending=False).reset_index(drop=True)


def build_top_keyword_report(df: pd.DataFrame, metric: str, n: int = 10, min_clicks: int = 0) -> pd.DataFrame:
    """상위 키워드 TOP N."""
    base = build_keyword_report(df)
    if base.empty:
        return base

    if min_clicks > 0:
        base = base[base["클릭수"] >= min_clicks]

    sort_map = {
        "노출": "노출수",
        "클릭": "클릭수",
        "광고비": "광고비",
        "CPC": "CPC",
        "클릭률": "클릭률",
        "전환": "전환수",
    }
    col = sort_map.get(metric, "클릭수")
    ascending = metric == "CPC"  # CPC TOP은 낮은 순이 아니라 높은 순 - user template uses high CPC
    if metric == "CPC":
        ascending = False
    return base.sort_values(col, ascending=ascending).head(n).reset_index(drop=True)


def build_summary_table(bundle: ReportBundle) -> pd.DataFrame:
    """요약리포트 - 캠페인/광고그룹별 (CSV 캠페인 요약 우선, 없으면 키워드 집계)."""
    if bundle.has_campaigns():
        rows = []
        for _, r in bundle.campaigns.iterrows():
            label = r.get("캠페인", r.get("광고그룹", ""))
            rows.append({"구분": label, **_calc_metrics_row(r["노출수"], r["클릭수"], r["광고비"], r.get("전환수", 0))})
        result = pd.DataFrame(rows)
    elif bundle.has_keywords():
        df = bundle.keywords.copy()
        if df["광고그룹"].astype(str).str.len().gt(0).any():
            df = df.rename(columns={"광고그룹": "구분"})
            result = _aggregate(df, ["구분"], report_style=True)
        else:
            df = df.rename(columns={"캠페인": "구분"})
            result = _aggregate(df, ["구분"], report_style=True)
    else:
        return pd.DataFrame()

    total = _calc_metrics_row(result["노출수"].sum(), result["클릭수"].sum(), result["광고비"].sum(), result.get("전환수", pd.Series([0])).sum())
    total["구분"] = "합계"
    return pd.concat([result, pd.DataFrame([total])], ignore_index=True)


def build_metric_table(df: pd.DataFrame, label_col: str = "구분") -> pd.DataFrame:
    """최종 리포트 형식 표 (노출·클릭·클릭률·CPC·광고비)."""
    if df.empty:
        return pd.DataFrame()

    label = label_col if label_col in df.columns else "구분"
    out = pd.DataFrame()
    out[label] = df[label] if label in df.columns else df.get("구분", "")

    for src, dst in [("노출수", "노출"), ("클릭수", "클릭"), ("광고비", "광고비")]:
        out[dst] = df[src] if src in df.columns else 0

    out["클릭률"] = df["클릭률"] if "클릭률" in df.columns else (
        out["클릭"] / out["노출"].replace(0, pd.NA)
    ).fillna(0).round(6)
    out["CPC"] = df["평균CPC"] if "평균CPC" in df.columns else (
        out["광고비"] / out["클릭"].replace(0, pd.NA)
    ).fillna(0).round(2)

    if "전환수" in df.columns and df["전환수"].sum() > 0:
        out["전환수"] = df["전환수"]
        out["전환율"] = (out["전환수"] / out["클릭"].replace(0, pd.NA)).fillna(0).round(6)
        out["CPA"] = (out["광고비"] / out["전환수"].replace(0, pd.NA)).fillna(0).round(2)

    return out


def _growth_rate(current: float, previous: float) -> float | None:
    """전월 대비 증감율 = (당월 - 전월) / 전월"""
    if previous == 0:
        return None
    return round((current - previous) / previous, 6)


def build_mom_comparison(current: dict, previous: dict | None) -> pd.DataFrame:
    """전월 대비 광고 데이터 비교표."""
    if not previous:
        return pd.DataFrame()

    has_conv = current.get("has_conversions") or previous.get("has_conversions")

    def _row(label: str, src: dict) -> dict:
        row = {
            "집행기간": label,
            "노출": src.get("노출수", 0),
            "클릭": src.get("클릭수", 0),
            "클릭률": src.get("클릭률", 0),
            "CPC": src.get("CPC", 0),
            "광고비": src.get("광고비", 0),
        }
        if has_conv:
            row["전환수"] = src.get("전환수", 0)
            row["전환율"] = src.get("전환율", 0)
            row["CPA"] = src.get("CPA", 0)
        return row

    cur_label = current.get("period_label") or current.get("report_month", "당월")
    prev_label = previous.get("period_label") or previous.get("report_month", "전월")

    cur_row = _row(cur_label, current)
    prev_row = _row(prev_label, previous)

    change: dict = {"집행기간": "전월 대비 증감율"}
    for key, ckey, pkey in [
        ("노출", "노출수", "노출수"),
        ("클릭", "클릭수", "클릭수"),
        ("클릭률", "클릭률", "클릭률"),
        ("CPC", "CPC", "CPC"),
        ("광고비", "광고비", "광고비"),
    ]:
        rate = _growth_rate(current.get(ckey, 0), previous.get(pkey, 0))
        change[key] = rate if rate is not None else ""

    if has_conv:
        for key, ckey, pkey in [
            ("전환수", "전환수", "전환수"),
            ("전환율", "전환율", "전환율"),
            ("CPA", "CPA", "CPA"),
        ]:
            rate = _growth_rate(current.get(ckey, 0), previous.get(pkey, 0))
            change[key] = rate if rate is not None else ""

    return pd.DataFrame([cur_row, change, prev_row])


def build_daily_report_from_bundle(bundle: ReportBundle) -> pd.DataFrame:
    if not bundle.has_daily():
        return pd.DataFrame()
    df = bundle.daily.copy()
    if "일자" in df.columns and df["일자"].notna().any():
        df = df.sort_values("일자")
        df["구분"] = pd.to_datetime(df["일자"]).dt.strftime("%Y.%m.%d.")
    return build_metric_table(df, "구분").rename(columns={"구분": "일자별"})


def build_weekday_report_from_bundle(bundle: ReportBundle) -> pd.DataFrame:
    if bundle.has_weekday():
        df = bundle.weekday.copy()
        order = {d: i for i, d in enumerate(KOREAN_WEEKDAYS)}
        df["_o"] = df["요일"].map(order)
        df = df.sort_values("_o").drop(columns=["_o"])
        return build_metric_table(df, "요일")
    if bundle.has_daily() and bundle.daily["일자"].notna().any():
        return build_weekday_report(bundle.daily)
    return pd.DataFrame()


def build_weekly_report_from_bundle(bundle: ReportBundle) -> pd.DataFrame:
    if not bundle.has_weekly():
        return pd.DataFrame()
    df = bundle.weekly.copy()
    label = "주차" if "주차" in df.columns else "구분"
    return build_metric_table(df, label)


def build_region_report_from_bundle(bundle: ReportBundle, detail: bool = False) -> pd.DataFrame:
    df = bundle.region_detail if detail and not bundle.region_detail.empty else bundle.region
    if df.empty:
        return pd.DataFrame()
    return build_metric_table(df, "지역").sort_values("클릭", ascending=False).reset_index(drop=True)


def build_demographic_report_from_bundle(bundle: ReportBundle) -> pd.DataFrame:
    """성별·연령·기기를 하나의 표로 (섹션 구분 행 포함)."""
    sections: list[pd.DataFrame] = []
    for title, df, col in [
        ("성별", bundle.gender, "성별"),
        ("연령", bundle.age, "연령"),
        ("기기", bundle.device, "기기"),
    ]:
        if df.empty:
            continue
        table = build_metric_table(df, col)
        header = pd.DataFrame([{col: f"[{title}]"}])
        sections.append(header)
        sections.append(table)
    if not sections:
        return pd.DataFrame()
    return pd.concat(sections, ignore_index=True)


def build_ad_group_report(df: pd.DataFrame) -> pd.DataFrame:
    """광고그룹별 성과."""
    group_cols = [c for c in ["캠페인", "광고그룹"] if c in df.columns and df[c].astype(str).str.len().gt(0).any()]
    if not group_cols:
        group_cols = ["광고그룹"] if "광고그룹" in df.columns else []
    if not group_cols:
        return _aggregate(df, [])
    result = _aggregate(df, group_cols)
    return result.sort_values("광고비", ascending=False).reset_index(drop=True)


def build_daily_report(df: pd.DataFrame) -> pd.DataFrame:
    """일자별 성과."""
    if df["일자"].isna().all():
        return pd.DataFrame(columns=["일자", "요일", "노출수", "클릭수", "광고비"])

    work = _add_weekday(df)
    result = _aggregate(work, ["일자", "요일"])
    result["일자"] = result["일자"].dt.strftime("%Y-%m-%d")
    return result.sort_values("일자").reset_index(drop=True)


def build_weekday_report(df: pd.DataFrame) -> pd.DataFrame:
    """요일별 성과."""
    if df["일자"].isna().all():
        return pd.DataFrame(columns=["요일", "노출수", "클릭수", "광고비"])

    work = _add_weekday(df)
    result = _aggregate(work, ["요일"])

    order = {day: i for i, day in enumerate(KOREAN_WEEKDAYS)}
    result["_order"] = result["요일"].map(order)
    result = result.sort_values("_order").drop(columns=["_order"]).reset_index(drop=True)
    return result


def build_campaign_report(df: pd.DataFrame) -> pd.DataFrame:
    """캠페인별 성과."""
    if "캠페인" not in df.columns or df["캠페인"].astype(str).str.len().eq(0).all():
        return pd.DataFrame()
    result = _aggregate(df, ["캠페인"])
    return result.sort_values("광고비", ascending=False).reset_index(drop=True)


def build_adgroup_daily_report(df: pd.DataFrame) -> pd.DataFrame:
    """광고그룹 × 일자 상세."""
    if df["일자"].isna().all():
        return pd.DataFrame()

    work = _add_weekday(df)
    group_cols = [c for c in ["캠페인", "광고그룹", "일자", "요일"] if c in work.columns]
    result = _aggregate(work, group_cols)
    result["일자"] = pd.to_datetime(result["일자"]).dt.strftime("%Y-%m-%d")
    return result.sort_values(["광고그룹", "일자"]).reset_index(drop=True)


_COL = 3  # 기존 리포트와 동일하게 C열(3)부터 데이터 배치


def _bundle_has_conversions(bundle: ReportBundle) -> bool:
    for df in (
        bundle.campaigns,
        bundle.keywords,
        bundle.daily,
        bundle.gender,
        bundle.age,
        bundle.device,
    ):
        if not df.empty and "전환수" in df.columns and df["전환수"].sum() > 0:
            return True
    return False


def _format_ad_period(bundle: ReportBundle) -> str:
    if not bundle.period_start or not bundle.period_end:
        return ""
    s, e = bundle.period_start, bundle.period_end
    days = (e - s).days + 1
    return f"{s.strftime('%Y/%m/%d')} - {e.strftime('%m/%d')} (총{days}일)"


def _format_period_short(bundle: ReportBundle) -> str:
    if not bundle.period_start or not bundle.period_end:
        return ""
    s, e = bundle.period_start, bundle.period_end
    days = (e - s).days + 1
    return f"{s.strftime('%m/%d')} - {e.strftime('%m/%d')} (총{days}일)"


def _metric_headers(has_conv: bool) -> list[str]:
    if has_conv:
        return ["노출", "클릭", "클릭률", "CPC", "전환수", "전환율", "CPA", "광고비"]
    return ["노출", "클릭", "클릭률", "CPC", "광고비"]


def _summary_row_to_metrics(row: pd.Series, has_conv: bool) -> list:
    vals = [
        int(row.get("노출수", row.get("노출", 0)) or 0),
        int(row.get("클릭수", row.get("클릭", 0)) or 0),
        row.get("클릭률", 0),
        row.get("CPC", 0),
    ]
    if has_conv:
        vals.extend([
            int(row.get("전환수", 0) or 0),
            row.get("전환율", 0),
            row.get("CPA", 0),
        ])
    vals.append(int(row.get("광고비", 0) or 0))
    return vals


def _metric_table_row(row: pd.Series, has_conv: bool) -> list:
    vals = [
        int(row.get("노출", 0) or 0),
        int(row.get("클릭", 0) or 0),
        row.get("클릭률", 0),
        row.get("CPC", 0),
    ]
    if has_conv:
        vals.extend([
            int(row.get("전환수", 0) or 0),
            row.get("전환율", 0),
            row.get("CPA", 0),
        ])
    vals.append(int(row.get("광고비", 0) or 0))
    return vals


def _apply_cell_style(cell, is_header: bool = False, is_rate: bool = False):
    thin = Side(style="thin", color="CCCCCC")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    cell.alignment = Alignment(horizontal="center" if is_header else "right", vertical="center")
    if is_header:
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.font = Font(color="FFFFFF", bold=True)
    elif isinstance(cell.value, (int, float)) and not is_rate:
        if abs(cell.value) < 1 and cell.value != 0:
            cell.number_format = "0.000000"
        else:
            cell.number_format = "#,##0"


def _write_row(ws, row: int, values: list, header: bool = False, rate_cols: set[int] | None = None):
    rate_cols = rate_cols or set()
    for i, val in enumerate(values):
        cell = ws.cell(row=row, column=_COL + i, value=val)
        _apply_cell_style(cell, is_header=header, is_rate=(i in rate_cols))


def _write_sheet_title(ws, row: int, title: str) -> int:
    cell = ws.cell(row=row, column=_COL, value=title)
    cell.font = Font(bold=True, size=14, color="1F4E79")
    return row + 3


def _write_period_lines(ws, row: int, bundle: ReportBundle) -> int:
    period = _format_ad_period(bundle)
    if period:
        ws.cell(row=row, column=_COL, value=f"광고기간 : {period}")
        row += 1
    ws.cell(row=row, column=_COL, value=f"작성일자 : {datetime.now().strftime('%Y.%m.%d')}")
    return row + 2


def _write_section_header(ws, row: int, title: str) -> int:
    cell = ws.cell(row=row, column=_COL, value=f"▶ {title}")
    cell.font = Font(bold=True, color="1F4E79")
    return row + 2


def _write_metrics_table(
    ws,
    row: int,
    label_header: str,
    table: pd.DataFrame,
    label_col: str,
    has_conv: bool,
    sum_first: bool = False,
) -> int:
    """지표 표 작성. sum_first=True면 '합계' 행을 헤더 바로 아래에."""
    if table.empty:
        return row

    headers = [label_header] + _metric_headers(has_conv)
    rate_offset = len(headers) - (5 if has_conv else 2)
    rate_cols = {rate_offset, rate_offset + (2 if has_conv else 0)}

    _write_row(ws, row, headers, header=True)
    row += 1

    data = table.copy()
    if sum_first and label_col in data.columns:
        total_mask = data[label_col].astype(str).str.contains("합계", na=False)
        totals = data[total_mask]
        rest = data[~total_mask]
        data = pd.concat([totals, rest], ignore_index=True)

    for _, r in data.iterrows():
        label = r.get(label_col, "")
        vals = [label] + _metric_table_row(r, has_conv)
        _write_row(ws, row, vals, rate_cols=rate_cols)
        row += 1

    return row + 1


def _write_summary_campaign_table(ws, row: int, summary: pd.DataFrame, bundle: ReportBundle, has_conv: bool) -> int:
    if summary.empty:
        return row

    headers = ["구분"] + _metric_headers(has_conv)
    rate_offset = len(headers) - (5 if has_conv else 2)
    rate_cols = {rate_offset, rate_offset + (2 if has_conv else 0)}

    period_short = _format_period_short(bundle)
    if period_short:
        ws.cell(row=row, column=_COL, value=f"{period_short} | (vat 포함가)")
        row += 1

    _write_row(ws, row, headers, header=True)
    row += 1

    for _, r in summary.iterrows():
        vals = [r.get("구분", "")] + _summary_row_to_metrics(r, has_conv)
        _write_row(ws, row, vals, rate_cols=rate_cols)
        row += 1

    return row + 1


def _write_mom_table(ws, row: int, mom: pd.DataFrame, has_conv: bool) -> int:
    if mom.empty:
        return row

    headers = ["집행기간"] + _metric_headers(has_conv)
    rate_offset = len(headers) - (5 if has_conv else 2)
    rate_cols = {rate_offset, rate_offset + (2 if has_conv else 0)}

    _write_row(ws, row, headers, header=True)
    row += 1

    for _, r in mom.iterrows():
        vals = [r.get("집행기간", "")] + _metric_table_row(r, has_conv)
        _write_row(ws, row, vals, rate_cols=rate_cols)
        row += 1

    return row + 1


def _keyword_top_columns(df: pd.DataFrame, has_conv: bool) -> tuple[list[str], list[str]]:
    """TOP10 표에 쓸 (라벨 컬럼들, 지표 컬럼들)"""
    label_cols: list[str] = []
    for c in ["캠페인", "광고그룹", "키워드", "검색유형"]:
        if c in df.columns and df[c].astype(str).str.len().gt(0).any():
            label_cols.append(c)

    if not label_cols:
        label_cols = ["키워드"] if "키워드" in df.columns else ["구분"]

    display = []
    for c in label_cols:
        display.append("검색어" if c == "키워드" else c)

    metrics = _metric_headers(has_conv)
    return display, label_cols + metrics


def _write_keyword_full_table(ws, row: int, keywords: pd.DataFrame, has_conv: bool) -> int:
    """키워드별 상세 — 캠페인·광고그룹·검색어 + 지표."""
    if keywords.empty:
        return row

    label_cols: list[str] = []
    for c in ["캠페인", "광고그룹", "키워드", "검색유형"]:
        if c in keywords.columns and keywords[c].astype(str).str.len().gt(0).any():
            label_cols.append(c)

    headers = []
    for c in label_cols:
        headers.append("검색어" if c == "키워드" else c)
    headers.extend(_metric_headers(has_conv))

    _write_row(ws, row, headers, header=True)
    row += 1

    rate_start = len(label_cols) + 2
    rate_cols = {rate_start, rate_start + (4 if has_conv else 0)}

    for _, r in keywords.iterrows():
        vals = [r.get(c, "") for c in label_cols]
        metric_row = pd.Series({
            "노출": r.get("노출수", 0),
            "클릭": r.get("클릭수", 0),
            "클릭률": r.get("클릭률", 0),
            "CPC": r.get("CPC", 0),
            "전환수": r.get("전환수", 0),
            "전환율": r.get("전환율", 0),
            "CPA": r.get("CPA", 0),
            "광고비": r.get("광고비", 0),
        })
        vals.extend(_metric_table_row(metric_row, has_conv))
        _write_row(ws, row, vals, rate_cols=rate_cols)
        row += 1

    return row + 1


def _write_keyword_top_section(
    ws,
    row: int,
    title: str,
    df: pd.DataFrame,
    has_conv: bool,
) -> int:
    if df.empty:
        return row

    row = _write_section_header(ws, row, title)
    display_cols, src_cols = _keyword_top_columns(df, has_conv)
    _write_row(ws, row, display_cols, header=True)
    row += 1

    metric_start = len(display_cols) - len(_metric_headers(has_conv))
    label_src = src_cols[:metric_start]
    rate_cols = set(range(metric_start + 2, metric_start + len(_metric_headers(has_conv))))

    for _, r in df.iterrows():
        vals = []
        for c in label_src:
            vals.append(r.get(c, ""))
        metric_row = pd.Series({
            "노출": r.get("노출수", r.get("노출", 0)),
            "클릭": r.get("클릭수", r.get("클릭", 0)),
            "클릭률": r.get("클릭률", 0),
            "CPC": r.get("CPC", 0),
            "전환수": r.get("전환수", 0),
            "전환율": r.get("전환율", 0),
            "CPA": r.get("CPA", 0),
            "광고비": r.get("광고비", 0),
        })
        vals.extend(_metric_table_row(metric_row, has_conv))
        _write_row(ws, row, vals, rate_cols=rate_cols)
        row += 1

    return row + 2


def _autofit_columns(ws, max_col: int = 12):
    for col_idx in range(_COL, _COL + max_col):
        letter = get_column_letter(col_idx)
        max_len = 10
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            val = row[0].value
            if val is not None:
                max_len = max(max_len, min(len(str(val)) + 2, 36))
        ws.column_dimensions[letter].width = max_len


def _build_workbook_sheets(
    wb,
    company_name: str,
    report_month: str,
    bundle: ReportBundle,
    cur: dict | None,
    mom: pd.DataFrame,
):
    has_conv = _bundle_has_conversions(bundle)
    account = f"({bundle.account_id})" if bundle.account_id else ""
    title_main = f"{company_name}{account} 네이버 검색광고 리포트"

    # 1. 요약리포트 (캠페인별 + 전월대비)
    ws = wb.create_sheet("요약리포트")
    row = _write_sheet_title(ws, 3, title_main)
    row = _write_period_lines(ws, row, bundle)
    row = _write_section_header(ws, row, "한달간 캠페인별 광고 데이터 총합")
    summary = build_summary_table(bundle)
    if not summary.empty:
        summary_display = build_metric_table(
            summary.rename(columns={"구분": "구분"}),
            "구분",
        )
        summary_display["구분"] = summary["구분"]
        row = _write_summary_campaign_table(ws, row, summary, bundle, has_conv)
    if not mom.empty:
        row = _write_section_header(ws, row, "전월 대비 광고 데이터 비교")
        row = _write_mom_table(ws, row, mom, has_conv)
    _autofit_columns(ws)

    # 2. 일자별&요일별
    daily = build_daily_report_from_bundle(bundle)
    weekday = build_weekday_report_from_bundle(bundle)
    if not daily.empty or not weekday.empty:
        ws = wb.create_sheet("일자별&요일별")
        row = _write_sheet_title(ws, 3, "일자별 및 요일별 상세 리포트")
        row = _write_period_lines(ws, row, bundle)
        if not daily.empty:
            row = _write_section_header(ws, row, "일자별 데이터")
            daily_sum = daily.copy()
            if not daily_sum.empty:
                total = pd.Series({
                    "일자별": "합계",
                    "노출": daily["노출"].sum(),
                    "클릭": daily["클릭"].sum(),
                    "클릭률": daily["클릭"].sum() / daily["노출"].sum() if daily["노출"].sum() else 0,
                    "CPC": daily["광고비"].sum() / daily["클릭"].sum() if daily["클릭"].sum() else 0,
                    "광고비": daily["광고비"].sum(),
                })
                if has_conv and "전환수" in daily.columns:
                    total["전환수"] = daily["전환수"].sum()
                    total["전환율"] = daily["전환수"].sum() / daily["클릭"].sum() if daily["클릭"].sum() else 0
                    total["CPA"] = daily["광고비"].sum() / daily["전환수"].sum() if daily["전환수"].sum() else 0
                daily_with_sum = pd.concat([pd.DataFrame([total]), daily], ignore_index=True)
            else:
                daily_with_sum = daily
            row = _write_metrics_table(ws, row, "일자별", daily_with_sum, "일자별", has_conv)
        if not weekday.empty:
            row = _write_section_header(ws, row, "요일별 데이터 (4주 총합)")
            row = _write_metrics_table(ws, row, "요일", weekday, "요일", has_conv)
        ws.cell(row=row + 1, column=_COL, value=f"{company_name} with us")
        _autofit_columns(ws)

    # 3. 주차별
    weekly = build_weekly_report_from_bundle(bundle)
    if not weekly.empty:
        ws = wb.create_sheet("주차별")
        row = _write_sheet_title(ws, 3, "주차별 상세 리포트")
        row = _write_period_lines(ws, row, bundle)
        label = "주차" if "주차" in weekly.columns else weekly.columns[0]
        row = _write_section_header(ws, row, "주차별 데이터")
        row = _write_metrics_table(ws, row, label, weekly, label, has_conv)
        _autofit_columns(ws)

    # 4. 키워드별
    keywords = build_keyword_report(bundle.keywords)
    if not keywords.empty:
        ws = wb.create_sheet("키워드별")
        row = _write_sheet_title(ws, 3, "키워드별 상세 리포트")
        row = _write_period_lines(ws, row, bundle)
        row = _write_section_header(ws, row, "키워드별 데이터 (클릭수 높은 순)")
        row = _write_keyword_full_table(ws, row, keywords, has_conv)
        _autofit_columns(ws)

    # 5. 상위 키워드 순위별 (TOP10 묶음)
    if bundle.has_keywords():
        ws = wb.create_sheet("상위 키워드 순위별")
        row = _write_sheet_title(ws, 3, "상위 키워드 순위별")
        row = _write_period_lines(ws, row, bundle)
        top_sections = [
            ("노출수 TOP 10", build_top_keyword_report(bundle.keywords, "노출")),
            ("클릭수 TOP 10", build_top_keyword_report(bundle.keywords, "클릭")),
        ]
        if has_conv:
            top_sections.append(("전환수 TOP 10", build_top_keyword_report(bundle.keywords, "전환")))
        top_sections.extend([
            ("광고비용 TOP 10", build_top_keyword_report(bundle.keywords, "광고비")),
            ("클릭당비용 TOP 10", build_top_keyword_report(bundle.keywords, "CPC")),
            ("클릭률 TOP 10 (2클릭이상)", build_top_keyword_report(bundle.keywords, "클릭률", min_clicks=2)),
        ])
        for title, section_df in top_sections:
            if section_df.empty:
                continue
            row = _write_keyword_top_section(ws, row, title, section_df, has_conv)
        _autofit_columns(ws)

    # 6. 지역별
    region = build_region_report_from_bundle(bundle, detail=False)
    if region.empty and not bundle.region_detail.empty:
        region = build_region_report_from_bundle(bundle, detail=True)
    if not region.empty:
        ws = wb.create_sheet("지역별")
        row = _write_sheet_title(ws, 3, "지역별 상세 리포트")
        row = _write_period_lines(ws, row, bundle)
        row = _write_section_header(ws, row, "지역별 클릭수 높은 순")
        label = "지역" if "지역" in region.columns else region.columns[0]
        row = _write_metrics_table(ws, row, label, region, label, has_conv)
        if not bundle.region_detail.empty and not bundle.region.empty:
            detail = build_region_report_from_bundle(bundle, detail=True)
            if not detail.empty:
                row = _write_section_header(ws, row, "상세지역별 클릭수 높은 순")
                label_d = "지역" if "지역" in detail.columns else detail.columns[0]
                row = _write_metrics_table(ws, row, label_d, detail, label_d, has_conv)
        _autofit_columns(ws)

    # 7. 성별&연령별&기기별
    if bundle.has_demographic():
        ws = wb.create_sheet("성별&연령별&기기별")
        row = _write_sheet_title(ws, 3, "성별/연령별/기기별 상세 리포트")
        row = _write_period_lines(ws, row, bundle)
        for section_title, df, col in [
            ("성별 상세", bundle.gender, "성별"),
            ("연령별", bundle.age, "연령"),
            ("기기별 상세", bundle.device, "기기"),
        ]:
            if df.empty:
                continue
            table = build_metric_table(df, col)
            row = _write_section_header(ws, row, section_title)
            row = _write_metrics_table(ws, row, col, table, col, has_conv)
        ws.cell(row=row + 1, column=_COL, value=f"{company_name} with us")
        _autofit_columns(ws)


def generate_excel_report(
    df: pd.DataFrame,
    company_name: str = "",
    report_month: str = "",
    bundle: ReportBundle | None = None,
    previous_summary: dict | None = None,
    current_metrics: dict | None = None,
) -> bytes:
    """기존 7시트 형식의 엑셀 리포트 생성 (시트 내 ▶ 섹션 묶음)."""
    from openpyxl import Workbook

    from .metrics import extract_monthly_metrics

    buffer = io.BytesIO()
    b = bundle or ReportBundle(keywords=df if not df.empty else pd.DataFrame())
    cur = current_metrics or (extract_monthly_metrics(b, report_month) if b.has_any_data() else None)
    mom = build_mom_comparison(cur, previous_summary) if cur and previous_summary else pd.DataFrame()

    wb = Workbook()
    wb.remove(wb.active)
    _build_workbook_sheets(wb, company_name, report_month, b, cur, mom)

    if not wb.sheetnames:
        ws = wb.create_sheet("요약리포트")
        ws.cell(row=3, column=_COL, value="데이터 없음")

    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def infer_report_month(bundle: ReportBundle | pd.DataFrame) -> str:
    """리포트 월 추론."""
    if isinstance(bundle, ReportBundle):
        if bundle.period_start:
            return bundle.period_start.strftime("%Y-%m")
        df = bundle.primary_df
    else:
        df = bundle

    if "일자" in df.columns and df["일자"].notna().any():
        d = df["일자"].dropna().iloc[0]
        return pd.Timestamp(d).strftime("%Y-%m")
    return datetime.now().strftime("%Y-%m")
